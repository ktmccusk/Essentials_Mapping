"""
utils/excel_writer.py
Writes AI-generated mapping JSON into the DNP curricular mapping Excel template.
"""

from __future__ import annotations
import logging
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

log = logging.getLogger(__name__)

# ── Cell styles ──────────────────────────────────────────────────────────────
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_IRD_FILL    = PatternFill("solid", fgColor="D9E1F2")
_EVAL_FILL   = PatternFill("solid", fgColor="E2EFDA")
_FLAG_FILL   = PatternFill("solid", fgColor="FFE699")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=9)
_BOLD_FONT   = Font(bold=True, size=10)
_DATA_FONT   = Font(size=9)
_CENTER      = Alignment(horizontal="center", vertical="center", wrap_text=True)
_THIN        = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)


def build_subcomp_index(wb: Workbook) -> dict[tuple[str, str], int]:
    """Map (sheet_name, subcomp_id) → row number, scanning column 3 from row 5."""
    index: dict[tuple[str, str], int] = {}
    for sname in wb.sheetnames:
        ws = wb[sname]
        for row in ws.iter_rows(min_row=5):
            cell = ws.cell(row=row[0].row, column=3)
            if cell.value and isinstance(cell.value, str):
                key = cell.value.strip()
                if key and key != "Subcomp #":
                    index[(sname, key)] = row[0].row
    return index


def _sheet_for(subcomp_id: str, framework: str) -> Optional[str]:
    sid = subcomp_id.strip()
    if framework.upper() == "NONPF" or sid.upper().startswith("NP "):
        try:
            num = sid.replace("NP ", "").replace("NP", "").split(".")[0].strip()
            return f"NONPF Domain {num}"
        except Exception:
            return None
    try:
        num = sid.split(".")[0].strip()
        return f"Domain {num}"
    except Exception:
        return None


def _find_or_create_col(ws, course_name: str) -> int:
    """
    Return the odd column (≥5) where this course's data lives in row 3.
    Creates the column header if the slot is empty.
    """
    col = 5
    while True:
        cell = ws.cell(row=3, column=col)
        val  = cell.value
        if val and course_name.lower() in str(val).lower():
            return col
        if val is None:
            cell.value     = course_name
            cell.font      = _HEADER_FONT
            cell.fill      = _HEADER_FILL
            cell.alignment = _CENTER
            ws.column_dimensions[get_column_letter(col)].width     = 16
            ws.column_dimensions[get_column_letter(col + 1)].width = 14
            return col
        col += 2


def write_courses(
    template_path: Path,
    output_path:   Path,
    mappings:      list[dict],
    subcomp_index: Optional[dict] = None,
) -> list[str]:
    """
    Write a list of course mapping dicts into the template and save to output_path.
    Returns a list of warning strings.
    """
    wb = load_workbook(str(template_path))
    if subcomp_index is None:
        subcomp_index = build_subcomp_index(wb)

    all_warnings: list[str] = []

    for course_map in mappings:
        course_name = course_map.get("course_name", "Unknown")
        course_num  = course_map.get("course_number", "")
        display     = f"{course_num} {course_name}".strip() if course_num else course_name
        col_cache: dict[str, int] = {}

        for entry in course_map.get("mappings", []):
            warns = _write_entry(wb, entry, subcomp_index, col_cache, display)
            all_warnings.extend(warns)

    wb.save(str(output_path))
    log.info("Saved: %s", output_path)
    return all_warnings


def _write_entry(
    wb:            Workbook,
    entry:         dict,
    subcomp_index: dict,
    col_cache:     dict[str, int],
    display_name:  str,
) -> list[str]:
    warnings: list[str] = []
    sid      = entry.get("subcomp_id", "").strip()
    framework= entry.get("framework", "AACN")
    ird      = entry.get("ird_code", "")
    eval_str = ", ".join(entry.get("eval_codes", []))
    flagged  = entry.get("flag") == "verify"

    sheet_name = _sheet_for(sid, framework)
    if not sheet_name or sheet_name not in wb.sheetnames:
        warnings.append(f"[{display_name}] No sheet for '{sid}' ({framework})")
        return warnings

    ws = wb[sheet_name]

    row_num = subcomp_index.get((sheet_name, sid))
    if not row_num:
        # Tolerate trailing whitespace variants in the template
        row_num = next(
            (r for (s, k), r in subcomp_index.items()
             if s == sheet_name and k.strip() == sid),
            None,
        )
    if not row_num:
        warnings.append(f"[{display_name}] Subcomp '{sid}' not found in '{sheet_name}'")
        return warnings

    if sheet_name not in col_cache:
        col_cache[sheet_name] = _find_or_create_col(ws, display_name)
    col = col_cache[sheet_name]

    fill = _FLAG_FILL if flagged else _IRD_FILL
    c1 = ws.cell(row=row_num, column=col)
    c1.value = ird; c1.font = _BOLD_FONT; c1.alignment = _CENTER
    c1.fill  = fill; c1.border = _THIN

    fill = _FLAG_FILL if flagged else _EVAL_FILL
    c2 = ws.cell(row=row_num, column=col + 1)
    c2.value = eval_str; c2.font = _DATA_FONT; c2.alignment = _CENTER
    c2.fill  = fill; c2.border = _THIN

    return warnings
